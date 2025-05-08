from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from datetime import datetime


class DocxExporter:
    @staticmethod
    def export(resume, filename):
        """Экспортирует резюме в DOCX файл"""
        doc = Document()

        # Настройка стилей
        style = doc.styles['Normal']
        font = style.font
        font.name = 'Arial'
        font.size = Pt(10)

        # Заголовок
        heading = doc.add_paragraph()
        heading.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
        runner = heading.add_run(resume.personal_info.get('name', ''))
        runner.bold = True
        runner.font.size = Pt(14)

        # Контактная информация
        contact = doc.add_paragraph()
        contact.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
        contact.add_run(f"{resume.personal_info.get('email', '')} | {resume.personal_info.get('phone', '')}")

        # Разделитель
        doc.add_paragraph().add_run().add_break()

        # Опыт работы
        doc.add_heading('Опыт работы', level=1)
        for exp in resume.work_experience:
            p = doc.add_paragraph()
            p.add_run(f"{exp.position} в {exp.company}").bold = True

            date_str = f"{exp.start_date.strftime('%m.%Y')} - {exp.end_date.strftime('%m.%Y') if exp.end_date else 'н.в.'}"
            p.add_run(f" ({date_str})")

            if exp.responsibilities:
                doc.add_paragraph('Обязанности:', style='List Bullet')
                for resp in exp.responsibilities:
                    doc.add_paragraph(resp, style='List Bullet 2')

            if exp.achievements:
                doc.add_paragraph('Достижения:', style='List Bullet')
                for ach in exp.achievements:
                    doc.add_paragraph(ach, style='List Bullet 2')

        # Образование
        doc.add_heading('Образование', level=1)
        for edu in resume.education:
            p = doc.add_paragraph()
            p.add_run(f"{edu.institution}").bold = True
            p.add_run(f", {edu.degree} по {edu.field_of_study}")

            date_str = f"{edu.start_date.strftime('%m.%Y')} - {edu.end_date.strftime('%m.%Y')}"
            doc.add_paragraph(date_str, style='List Bullet')

        # Навыки
        doc.add_heading('Навыки', level=1)
        skills_by_category = {}
        for skill in resume.skills:
            if skill.category not in skills_by_category:
                skills_by_category[skill.category] = []
            skills_by_category[skill.category].append(f"{skill.name} ({'★' * skill.level})")

        for category, skills in skills_by_category.items():
            doc.add_paragraph(category + ':', style='List Bullet')
            for skill in skills:
                doc.add_paragraph(skill, style='List Bullet 2')

        # Сохранение документа
        doc.save(filename)